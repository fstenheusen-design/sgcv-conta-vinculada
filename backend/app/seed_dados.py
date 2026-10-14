"""Carga inicial a partir da planilha de referência.

Coloque o arquivo conta_vinculada_detran_ro.xlsx em backend/dados/ antes de rodar.
Cabeçalhos reais da planilha (aba Contratos): Nº Contrato, Contratada, Objeto,
Vigência Início, Vigência Fim, Faturamento Mensal (R$).
Cabeçalhos reais (aba Empregados): Nº Contrato, Matrícula, Nome, Função,
Salário Base (R$), Data Admissão.
"""
import unicodedata
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.database import SessaoLocal
from app.models.contrato import Contrato
from app.models.empregado import Empregado
from app.models.parametro import ParametroVigencia

ARQUIVO = Path(__file__).resolve().parent.parent / "dados" / "conta_vinculada_detran_ro.xlsx"

INSS_PADRAO = Decimal("0.20")
RAT_PADRAO = Decimal("0.03")
FGTS_PADRAO = Decimal("0.08")
MULTA_PADRAO = Decimal("0.50")


def _norm(texto) -> str:
    """Minúsculas, sem acentos — para casar cabeçalhos como 'Nº Contrato'."""
    s = str(texto).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def _linhas(ws):
    linhas = list(ws.iter_rows(values_only=True))
    if len(linhas) < 2:
        return []
    cab = [_norm(c) if c is not None else "" for c in linhas[0]]
    dados = []
    for linha in linhas[1:]:
        if all(c is None for c in linha):
            continue
        dados.append({cab[i]: linha[i] for i in range(len(cab)) if cab[i]})
    return dados


def _pegar(linha: dict, *chaves: str):
    """Retorna o valor da primeira chave que casar (busca por substring)."""
    for chave in chaves:
        for nome, valor in linha.items():
            if chave in nome:
                return valor
    return None


def _numero(v) -> str:
    """Normaliza o nº do contrato (pode vir como número ou texto)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _dec(v) -> Decimal:
    if v is None:
        return Decimal("0")
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    texto = str(v).replace("R$", "").replace(" ", "").strip()
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto or "0")
    except Exception:
        return Decimal("0")


def _data(v, padrao):
    if hasattr(v, "year"):
        return date(v.year, v.month, v.day)
    return padrao


def main():
    if not ARQUIVO.exists():
        print(f"ERRO: planilha não encontrada em {ARQUIVO}")
        print("Copie o arquivo conta_vinculada_detran_ro.xlsx para backend/dados/")
        return

    wb = load_workbook(ARQUIVO, data_only=True)
    print(f"Abas encontradas: {wb.sheetnames}")

    ws_contratos = wb["Contratos"] if "Contratos" in wb.sheetnames else wb[wb.sheetnames[1]]
    ws_empregados = wb["Empregados"] if "Empregados" in wb.sheetnames else wb[wb.sheetnames[2]]

    db = SessaoLocal()
    try:
        # 1. Contratos + parâmetros padrão (usa a Vigência Início da planilha)
        mapa_contratos: dict[str, Contrato] = {}
        for linha in _linhas(ws_contratos):
            numero = _numero(_pegar(linha, "contrato"))
            contratada = str(_pegar(linha, "contratada") or "").strip()
            if not numero or not contratada:
                continue
            existente = db.query(Contrato).filter(Contrato.numero == numero).first()
            if existente:
                print(f"Contrato {numero} já existe — usando o existente")
                mapa_contratos[numero] = existente
                continue
            vig_inicio = _data(_pegar(linha, "vigencia inicio"), date(2026, 1, 1))
            novo = Contrato(
                numero=numero,
                contratada=contratada,
                faturamento_mensal=_dec(_pegar(linha, "faturamento")),
            )
            db.add(novo)
            db.flush()
            mapa_contratos[numero] = novo
            print(f"Contrato criado: {numero} — {contratada}")
            db.add(ParametroVigencia(
                contrato_id=novo.id,
                vigencia_inicio=vig_inicio,
                inss_patronal=INSS_PADRAO,
                rat=RAT_PADRAO,
                fgts=FGTS_PADRAO,
                multa_rescisoria=MULTA_PADRAO,
            ))
            print(f"  Parâmetros padrão criados (vigência a partir de {vig_inicio})")

        if not mapa_contratos:
            print("Nenhum contrato encontrado na aba 'Contratos'.")
            return

        # 2. Empregados vinculados pelo Nº Contrato (coluna presente na planilha)
        criados = 0
        for linha in _linhas(ws_empregados):
            numero = _numero(_pegar(linha, "contrato"))
            matricula = str(_pegar(linha, "matricula") or "").strip()
            nome = str(_pegar(linha, "nome") or "").strip()
            if not matricula or not nome:
                continue
            contrato = mapa_contratos.get(numero)
            if contrato is None:
                print(f"AVISO: empregado {matricula} refere o contrato '{numero}', não cadastrado — pulando")
                continue
            if db.query(Empregado).filter(
                Empregado.contrato_id == contrato.id, Empregado.matricula == matricula
            ).first():
                continue
            db.add(Empregado(
                contrato_id=contrato.id,
                matricula=matricula,
                nome=nome,
                funcao=str(_pegar(linha, "funcao") or "Não informado").strip(),
                salario_base=_dec(_pegar(linha, "salario")),
                data_admissao=_data(_pegar(linha, "admissao"), date(2025, 1, 1)),
            ))
            criados += 1

        db.commit()
        print(f"\nCarga concluída! {criados} empregados cadastrados.")
        print(f"Contratos no banco: {', '.join(mapa_contratos)}")
    finally:
        db.close()


if __name__ == "__main__":
    main()